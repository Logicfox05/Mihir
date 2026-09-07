from __future__ import annotations

import io


def _rows_to_dicts(rows: list[list]) -> list[dict]:
    rows = [list(r) for r in rows]
    while rows and not any(c is not None and str(c).strip() for c in rows[0]):
        rows.pop(0)
    if not rows:
        return []
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    out: list[dict] = []
    for r in rows[1:]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        out.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    return out


def parse(body: bytes, sheet: str | None = None) -> list[dict]:
    if body[:2] == b"PK":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return _rows_to_dicts(rows)
    # legacy .xls
    import xlrd

    book = xlrd.open_workbook(file_contents=body)
    sh = book.sheet_by_name(sheet) if sheet and sheet in book.sheet_names() else book.sheet_by_index(0)
    rows = [sh.row_values(i) for i in range(sh.nrows)]
    return _rows_to_dicts(rows)
